"""DOCX text extraction using python-docx.

Pulls paragraphs *and* table cells, because a large share of real resumes lay
their skills out in invisible tables.
"""

from __future__ import annotations

import io
import logging

from app.models.models import ParsedDocument

logger = logging.getLogger(__name__)


def extract_docx(data: bytes, filename: str = "document.docx") -> ParsedDocument:
    """Extract text from DOCX bytes."""
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        from app.parsers.pdf_parser import DocumentParseError

        raise DocumentParseError(
            "python-docx is not installed. Run: pip install -r requirements.txt"
        ) from exc

    if not data:
        return ParsedDocument(
            filename=filename,
            text="",
            file_type="docx",
            extraction_ok=False,
            warnings=["The uploaded DOCX is empty (0 bytes)."],
        )

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        logger.warning("Failed to open DOCX %s: %s", filename, exc)
        return ParsedDocument(
            filename=filename,
            text="",
            file_type="docx",
            extraction_ok=False,
            warnings=[
                "The DOCX could not be opened - it may be corrupt or actually be "
                f"an older .doc file. ({exc})"
            ],
        )

    lines: list[str] = []
    try:
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                cells = [c for c in cells if c]
                # De-duplicate merged cells that repeat the same text.
                deduped: list[str] = []
                for cell in cells:
                    if not deduped or deduped[-1] != cell:
                        deduped.append(cell)
                if deduped:
                    lines.append(" | ".join(deduped))
    except Exception as exc:
        logger.warning("Error while reading DOCX body %s: %s", filename, exc)
        return ParsedDocument(
            filename=filename,
            text="\n".join(lines),
            file_type="docx",
            extraction_ok=bool(lines),
            warnings=[f"Part of the document could not be read. ({exc})"],
        )

    text = "\n".join(lines).strip()
    if not text:
        return ParsedDocument(
            filename=filename,
            text="",
            file_type="docx",
            extraction_ok=False,
            warnings=[
                "No readable text was found in the DOCX. If the content is inside "
                "images or text boxes, please upload a PDF or plain text version."
            ],
        )

    return ParsedDocument(
        filename=filename,
        text=text,
        char_count=len(text),
        file_type="docx",
        extraction_ok=True,
    )
