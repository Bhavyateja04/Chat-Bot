"""Generate sample PDF and DOCX files for the Discord demo.

The test fixtures are plain text so the suite stays readable and fast. A live
demo is more convincing with real PDFs and DOCX files to drag into Discord, so
this script renders the fixtures into `samples/`.

    python make_samples.py
"""

from __future__ import annotations

import pathlib
import sys

FIXTURES = pathlib.Path(__file__).parent / "tests" / "fixtures"
SAMPLES = pathlib.Path(__file__).parent / "samples"

# (fixture, output basename, format)
TARGETS: list[tuple[str, str, str]] = [
    ("software_jd.txt", "job_description_software_engineer", "pdf"),
    ("mechanical_jd.txt", "job_description_mechanical_engineer", "pdf"),
    ("ml_engineer_jd.txt", "job_description_ml_engineer", "pdf"),
    ("backend_jd.txt", "job_description_backend_engineer", "txt"),
    ("strong_software_resume.txt", "resume_priya_raghavan_strong", "pdf"),
    ("weak_software_resume.txt", "resume_arjun_mehta_junior", "pdf"),
    ("fullstack_resume.txt", "resume_neha_kulkarni_fullstack", "docx"),
    ("mechanical_resume.txt", "resume_rohit_deshmukh_mechanical", "pdf"),
]


def write_pdf(lines: list[str], destination: pathlib.Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    width, height = A4
    pdf = canvas.Canvas(str(destination), pagesize=A4)
    margin, y = 56, height - 60

    for line in lines:
        if y < 60:
            pdf.showPage()
            y = height - 60

        stripped = line.strip()
        if not stripped:
            y -= 8
            continue

        # Headings (short lines with no leading bullet) get a bold face.
        is_heading = (
            len(stripped) < 60
            and not stripped.startswith(("-", "•"))
            and not stripped.endswith(".")
        )
        pdf.setFont("Helvetica-Bold" if is_heading else "Helvetica", 11 if is_heading else 9.5)

        # Wrap long lines rather than letting them run off the page.
        max_chars = 95 if not is_heading else 70
        while stripped:
            chunk, stripped = stripped[:max_chars], stripped[max_chars:]
            if stripped and not stripped[0].isspace() and " " in chunk:
                cut = chunk.rfind(" ")
                chunk, stripped = chunk[:cut], chunk[cut:] + stripped
            pdf.drawString(margin, y, chunk.strip())
            y -= 14 if is_heading else 12
            if y < 60:
                pdf.showPage()
                y = height - 60
        y -= 3 if is_heading else 0

    pdf.save()


def write_docx(lines: list[str], destination: pathlib.Path) -> None:
    import docx

    document = docx.Document()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped) < 60 and not stripped.startswith(("-", "•")) and not stripped.endswith("."):
            document.add_heading(stripped, level=2)
        else:
            document.add_paragraph(stripped)
    document.save(str(destination))


def main() -> int:
    if not FIXTURES.exists():
        print(f"Fixtures not found at {FIXTURES}", file=sys.stderr)
        return 1

    SAMPLES.mkdir(exist_ok=True)
    for fixture_name, basename, fmt in TARGETS:
        source = FIXTURES / fixture_name
        if not source.exists():
            print(f"  skipped (missing): {fixture_name}")
            continue

        lines = source.read_text(encoding="utf-8").splitlines()
        destination = SAMPLES / f"{basename}.{fmt}"

        if fmt == "pdf":
            write_pdf(lines, destination)
        elif fmt == "docx":
            write_docx(lines, destination)
        else:
            destination.write_text("\n".join(lines), encoding="utf-8")

        print(f"  {destination.relative_to(SAMPLES.parent)}  ({destination.stat().st_size:,} bytes)")

    print(f"\nSample files written to {SAMPLES}")
    print("Drag these straight into Discord to demo the bot.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
